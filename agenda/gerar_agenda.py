#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gerar_agenda.py
Preenche a agenda do site MNS4A (index.html) a partir da planilha do formulário de inscrição.

Uso
---
    python gerar_agenda.py MNS4A_Offline_com_agenda.xlsx
        Atualiza o index.html (guarda uma cópia em index.html.bak).

    python gerar_agenda.py planilha.xlsx --saida index_teste.html
        Não mexe no index.html; escreve outra página para conferir.

    python gerar_agenda.py planilha.xlsx --saida index_teste.html --ordem-das-linhas
        PRÉVIA: quem ainda não tem horário entra na ordem das linhas da planilha.

    --aba "Session1,Session2,Session3"   lê várias abas (padrão: aba 'Programa' ou a primeira)
    --estrito                            se houver QUALQUER problema na planilha, não escreve nada
    --exemplo                            acrescenta o aviso "dados fictícios" (só para demonstrações)

O que o script lê da planilha (exportação do Google Forms, sem precisar reformatar)
------------------------------------------------------------------------------
    Presentation Type · Presentation title · Authors · Abstract · Authors Institutions ·
    Full Name (apresentador) · Thematic Session · Timestamp (vira a chave de cada trabalho)
    E-mail, crachá, país e inscrições sem apresentação são ignorados.

    Colunas que você acrescenta para montar a agenda (todas opcionais):
      Day         1, 2 ou 3. Em branco = o dia da sessão temática (SESSION 01 = dia 1...).
      Order       posição no dia entre TODOS os horários, em ordem de relógio (1 = 09:00 convidado,
                  2 = 09:40 oral... até 11). Volta a 1 em cada dia. Convidado só cabe nos horários de convidado.
      Start time  horário exato (09:40). Vale mais que o Order.
    Sem Order nem Start time, o trabalho fica fora da agenda (o horário continua TBD).
    Order maior que o número de horários do dia (11) também fica de fora (com aviso), sem bloquear o resto.
    Convidados: linha extra com Presentation Type = Invited Talk (título, autores, abstract, Start time).

    Pôsteres (Presentation Type = Poster): ficam expostos todos os dias, então não têm dia nem horário.
      Cada pôster entra no bloco do seu TEMA (Thematic Session: Modeling, Networks...). A seção "Poster session"
      é montada pelo script: um bloco por tema que tem pôster (na ordem dos dias da agenda), um pôster por
      linha, empilhados. Tema sem nenhum pôster não aparece (hoje, Instrumentation); se alguém cadastrar um
      pôster de um tema que ainda não tinha bloco, o bloco é criado e o script avisa.
      O número (P01, P02...) é gerado pelo script, em sequência contínua e sem lacunas, seguindo a ordem
      dos blocos da página. A coluna Poster (ou Order), se preenchida ("P07", "p7" ou 7), só decide a
      ORDEM dentro do tema; em branco, vale a ordem das linhas da planilha.
      Slots reservados: um cartão numerado em TBD que você deixa no bloco da página
      (<div class="poster-card" data-poster="P12">…<span class="poster-content tbd">TBD</span>…) é
      mantido pelo script, depois dos pôsteres do bloco, e some sozinho quando um pôster NOVO da planilha
      o ocupa (a numeração continua contínua). Para criar um slot, copie uma linha TBD no HTML; para
      tirar, apague a linha. O aviso "To be announced" de bloco vazio não conta como slot.
      Página no formato ANTIGO (3 colunas, Day 01/02/03): o script migra sozinho para o formato novo (blocos,
      título, parágrafo e CSS dos pôsteres) e descarta os cartões TBD antigos. Depois da migração, o título e o
      parágrafo passam a ser seus: o script não mexe mais neles.

O que o script faz
------------------
* Autores em qualquer dos formatos usados no formulário: "[1] Nome; [2] Nome", "Nome[1;2]",
  "[1] Nome, [2] Nome", "Nome; Nome"... Mostra o primeiro autor + "et al." no cartão; a janela do
  abstract mostra todos os autores com as instituições numeradas.
* Ao clicar no título, abre uma janela com título, autores, instituições e abstract.
* Confere a planilha (horário ocupado duas vezes, horário que não existe, número de pôster inválido,
  tipo errado...). Linha com problema fica FORA da página e é listada no fim; o resto é escrito
  (com --estrito, nada é escrito enquanto houver problema). Avisa o que vale revisar (abstract de teste,
  apresentador diferente do primeiro autor, mais orais do que horários, pôster de um tema que não
  tem bloco na página...).
* Pode ser rodado quantas vezes quiser: o resultado é sempre o mesmo. Com a planilha sem nenhum
  trabalho escalado, a agenda volta ao original (horários TBD) e cada bloco de pôsteres mostra só
  os slots reservados ou, se não houver nenhum, "To be announced".
* A estrutura de horários (intervalos, almoço, mesa-redonda) continua sendo a do HTML.
* Os blocos da seção "Poster session" são refeitos inteiros a cada execução, a partir da planilha: um cartão
  por pôster (título + primeiro autor "et al."; clicar abre o abstract), sem cartões vazios. Os pôsteres não têm dia nem horário;
  a janela do abstract mostra só o tema e o número (ex.: "Modeling · Poster P01").
* A frase de dica da agenda ("Click a title to read its abstract.") fica entre os marcadores
  PROGRAMA:INTRO e é regravada a cada execução.

Requisitos: Python 3 e openpyxl  (pip install openpyxl)
"""

import argparse
import datetime as dt
import html
import json
import re
import shutil
import sys
import unicodedata
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Falta a biblioteca openpyxl. Instale com:  pip install openpyxl")


# ============================================================================
# Textos fixos
# ============================================================================
INTRO_ORIGINAL = (
    "The schedule below shows the time structure of each day; "
    "speakers and titles will be assigned after the call for contributions closes."
)
INTRO_PREENCHIDA = "Click a title to read its abstract."
AVISO_EXEMPLO = (
    '<div class="demo-banner">Example &mdash; fictitious data, layout preview only. '
    "Do not publish.</div>"
)
AVISO_PROVISORIO = (
    '<div class="demo-banner">Preview &mdash; provisional order (spreadsheet row order), not the final program. '
    "Do not publish.</div>"
)

# Nome da coluna na planilha (sem acento, minúsculo, sem símbolos) -> campo interno.
# Vale o nome exato OU, nos apelidos com 8+ letras, o começo do nome: assim os cabeçalhos longos
# do Google Forms ("Presentation title (max. 150 characters)") são reconhecidos.
CAMPOS = {
    "id": ["id", "codigo", "key", "chave"],
    "timestamp": ["timestamp", "carimbodedatahora"],
    "tipo": ["tipo", "type", "presentationtype", "tipodeapresentacao"],
    "sessao": ["thematicsession", "sessaotematica", "sessao", "session"],
    "dia": ["dia", "day"],
    "ordem": ["ordem", "order", "posicao"],
    "pnum": ["poster", "posternumber", "posterno", "numerodoposter", "numeropainel"],
    "inicio": ["inicio", "start", "starttime", "hora", "horario"],
    "fim": ["fim", "end", "endtime"],
    "titulo": ["titulo", "title", "presentationtitle", "titulodaapresentacao"],
    "apresentador": ["apresentador", "presenter", "speaker", "fullname", "nomecompleto"],
    "autores": ["autores", "autor", "authors", "authorsfullnames", "autoresnomescompletos"],
    "afiliacoes": ["afiliacoes", "afiliacao", "affiliations", "affiliation",
                   "authorsinstitutions", "instituicoesdosautores"],
    "abstract": ["abstract", "resumo"],
}
# Só estas colunas precisam existir. Todo o resto do formulário (e-mail, crachá, país...) é ignorado.
OBRIGATORIOS = ["tipo", "titulo", "autores", "abstract"]

PARTICULAS = {"de", "da", "do", "dos", "das", "e", "di", "du", "del", "della", "van", "von", "der", "den", "la", "le"}

DAY_RE = re.compile(r"<!--\s*DAY\s+(\d+)\s*-->")
SLOT_RE = re.compile(
    r'^(?P<indent>[ \t]*)<div class="slot(?P<cls>[^"]*)">'
    r'(?:<span class="slot-meta">)?<span class="slot-time">(?P<tempo>[^<]+)</span>.*</div>[ \t]*$'
)
TEMPO_RE = re.compile(r"(\d{1,2}:\d{2})\s*[–-]\s*(\d{1,2}:\d{2})")
# Seção de pôsteres da página: um bloco por tema (Modeling, Networks...), cada um com
#   <div class="poster-day-theme">Modeling</div>   e   <div class="poster-list"> ...cartões... </div>
# Os cartões (<div class="poster-card" data-poster="P01">...</div>, um por linha) são gerados pelo script.
POSTER_SESSION_RE = re.compile(r"<!--\s*POSTER SESSION\s*-->")
POSTER_TEMA_RE = re.compile(r'class="poster-day-theme">([^<]+)<')
POSTER_LISTA_RE = re.compile(r'<div class="poster-list">')
POSTER_CARD_RE = re.compile(r'^(?P<indent>[ \t]*)<div class="poster-card\b[^"]*"[^>]*>.*</div>[ \t]*$')
# Slot reservado: cartão numerado e ainda sem pôster (<div class="poster-card" data-poster="P12">…TBD…</div>).
# O aviso "To be announced" de bloco vazio não tem data-poster, então não conta como reserva.
SLOT_RESERVADO_RE = re.compile(r'data-poster="P?\d+"[^>]*>.*poster-content tbd')
POSTER_DAYS_RE = re.compile(r'<div class="poster-days">')
VERSAO = "2026-10-06 · seção de pôsteres por tema, um pôster por linha (migra o formato antigo de 3 colunas)"

# CSS da seção de pôsteres no formato novo. Só é gravado quando a página ainda está no formato antigo
# (3 colunas); em página já migrada o CSS não é tocado.
POSTER_CSS_NOVO = r'''  /* POSTER SESSION (abaixo da agenda)
     Classes próprias (poster-*), de propósito diferentes de .slot / .day-date / .day-theme,
     para o gerar_agenda.py não confundir estes cartões com os horários da agenda. */
  .poster-session { margin-top: 56px; }

  .poster-session h3 {
    font-family: 'Cormorant Garamond', serif;
    font-size: 30px;
    font-weight: 600;
    line-height: 1.2;
    color: var(--white);
    margin-bottom: 10px;
  }
  .poster-session h3 em { font-style: italic; color: var(--electric); }

  /* Blocos de tema empilhados (um em cima do outro), cada um com seu próprio contorno,
     para ninguém comparar as sessões lado a lado. Um pôster por linha. */
  .poster-days {
    display: flex;
    flex-direction: column;
    gap: 28px;
    margin-top: 24px;
  }

  .poster-day {
    background: var(--navy);
    border: 1px solid var(--border);
  }

  /* complemento do título: "(remain on display every day)", em etiqueta destacada (verde-água) */
  .poster-session h3 .poster-sub {
    display: inline-block;
    margin-left: 6px;
    padding: 3px 12px;
    font-family: 'DM Mono', monospace;
    font-size: 12px;
    font-style: normal;
    font-weight: 400;
    line-height: 1.5;
    letter-spacing: 0.06em;
    color: var(--aurora);
    background: rgba(168, 237, 204, 0.08);
    border: 1px solid rgba(168, 237, 204, 0.4);
    border-radius: 2px;
    vertical-align: middle;
    position: relative;
    top: -3px;
  }

  .poster-day-head {
    padding: 14px 20px 16px;
    background: var(--navy-mid);
    border-bottom: 1px solid var(--border);
  }
  .poster-day-theme {
    font-family: 'Cormorant Garamond', serif;
    font-size: 28px;
    font-weight: 600;
    font-style: italic;
    line-height: 1.1;
    color: var(--aurora);
  }

  .poster-list { display: flex; flex-direction: column; }

  .poster-card {
    display: grid;
    grid-template-columns: 48px 1fr;
    column-gap: 14px;
    align-items: center;
    box-sizing: border-box;
    min-height: 56px;
    padding: 12px 20px;
    border-bottom: 1px solid var(--border);
    transition: background 0.15s;
  }
  .poster-card:last-child { border-bottom: none; }
  .poster-card:hover { background: var(--navy-mid); }

  .poster-no {
    font-family: 'DM Mono', monospace;
    font-size: 12px;
    letter-spacing: 0.04em;
    color: var(--aurora-dim);
  }
  .poster-content { font-size: 13.5px; line-height: 1.4; color: #d8e8f8; min-width: 0; }
  .poster-content.tbd { color: var(--muted); }

  /* Cartão ocupa a largura toda: título completo, sem altura fixa e sem corte com "…". */
  @media (max-width: 480px) {
    .poster-card { grid-template-columns: 40px 1fr; column-gap: 10px; padding: 11px 14px; }
  }'''
FECHA_DIV_RE = re.compile(r"^[ \t]*</div>[ \t]*$")
LIGATURAS = {"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl"}


class ErroPlanilha(Exception):
    pass


# ============================================================================
# Leitura e limpeza da planilha
# ============================================================================
def norm(s):
    """Minúsculo, sem acento e sem símbolos: 'Início ' -> 'inicio'."""
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", s.lower())


def limpar(s):
    s = s.replace("\u00ad", "")                      # hífen opcional (colado de PDF)
    for a, b in LIGATURAS.items():                   # ligaduras (colado de PDF)
        s = s.replace(a, b)
    return s.replace("\r\n", "\n").replace("\r", "\n").strip()


def txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return limpar(str(v))


def hhmm(v):
    """Aceita texto '9:40', hora do Excel (time/datetime) ou fração de dia."""
    if v is None or v == "":
        return ""
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, (int, float)) and 0 <= v < 1:
        m = round(v * 24 * 60)
        return f"{m // 60:02d}:{m % 60:02d}"
    s = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})\s*[:hH.]\s*(\d{2})(?::\d{2})?", s)
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else s


def parse_dia(v):
    if isinstance(v, (dt.datetime, dt.date)):
        n = v.day
    else:
        m = re.search(r"\d+", txt(v))
        n = int(m.group()) if m else None
    if n in (1, 2, 3):
        return n
    if n in (19, 20, 21):                            # aceita também a data (19/20/21)
        return n - 18
    return None


BRACKET_RE = re.compile(r"\[([^\]]*)\]")
INICIAIS_RE = re.compile(r"^(?:[A-ZÀ-Ý]\.?\s*){1,4}$")
INSTITUICAO_RE = re.compile(r"institut|universi|laborat|observat|depart|facult|centre|center|agency|council", re.I)


def tipo_de(bruto):
    """'Oral Presentation' / 'Poster Presentation' / 'Invited talk' / 'convidado' ... -> invited | oral | poster."""
    n = norm(bruto)
    if not n:
        return None
    if "invit" in n or "convid" in n:
        return "invited"
    if "poster" in n or "painel" in n:
        return "poster"
    if n.startswith("oral") or n in ("talk", "palestra", "apresentacaooral"):
        return "oral"
    return None


def dia_da_sessao(v):
    """'SESSION 02 · NETWORKS · ...' -> 2  (a sessão temática de cada dia)."""
    m = re.search(r"session\w*\s*0*(\d)|sess[aã]o\s*0*(\d)", txt(v), re.I)
    n = int(next(g for g in m.groups() if g)) if m else None
    return n if n in (1, 2, 3) else None


def ajusta_caixa(nome):
    """Só mexe em nomes escritos TODOS EM MAIÚSCULAS: 'JOAO DA SILVA' -> 'Joao da Silva'."""
    if len(nome) > 3 and nome == nome.upper() and any(c.isalpha() for c in nome):
        palavras = [w.lower() if w.lower() in PARTICULAS else w.capitalize() for w in nome.lower().split()]
        return " ".join(palavras)
    return nome


def limpa_nome(seg):
    seg = re.sub(r"\s+", " ", seg).strip(" \t;,:")
    seg = re.sub(r"(?i)[\s,]+et\.? ?al\.?$", "", seg)          # "C. Dias et al." -> "C. Dias"
    if not any(c.isalpha() for c in seg) or re.fullmatch(r"(?i)et\.? ?al\.?", seg):
        return ""
    return ajusta_caixa(seg)


def marcas_de(conteudo):
    return [p for p in re.split(r"[,;\s]+", conteudo.strip()) if p]


def parse_autores(v):
    """
    Entende os formatos que aparecem no formulário:
        [1] Ana Silva; [2] Bruno Costa        (número antes do nome)
        R. de Jesus[1;2]; G. Yang[1]          (número depois do nome)
        [1] A. Silva, [2] B. Costa, C. Dias   (separado por vírgula)
        Ana Silva; Bruno Costa                (sem números)
    Devolve ([{'n': nome, 'm': [marcas]}], [avisos]).
    """
    s = txt(v)
    avisos = []
    if not s:
        return [], avisos
    sem_marcas = BRACKET_RE.sub(" ", s)
    tem_pv = ";" in sem_marcas
    usa_virgula = not tem_pv                                # vírgula só separa autores se não houver ';'
    if tem_pv or BRACKET_RE.search(s):
        s = s.replace("\n", " ")                            # Enter no meio do nome ("Giorgio⏎Picanço") não separa autores
        sep = r"[,;:]+" if usa_virgula else r"[;:]+"
    else:
        sep = r"[,;\n:]+"                                   # sem ';' nem números: um autor por linha também vale

    tokens, pos = [], 0
    for m in BRACKET_RE.finditer(s):
        if m.start() > pos:
            tokens.append(("t", s[pos:m.start()]))
        tokens.append(("m", marcas_de(m.group(1))))
        pos = m.end()
    if pos < len(s):
        tokens.append(("t", s[pos:]))
    prefixo = bool(tokens) and tokens[0][0] == "m"

    autores, pend = [], []
    for kind, val in tokens:
        if kind == "m":
            if prefixo:
                pend = val
            elif autores:
                autores[-1]["m"] += [x for x in val if x not in autores[-1]["m"]]
            continue
        for i, seg in enumerate(re.split(sep, val)):
            nome = limpa_nome(seg)
            if not nome:
                continue
            if usa_virgula and i > 0 and autores and INICIAIS_RE.match(nome):
                autores[-1]["n"] += ", " + nome                # "Silva, J." separado pela vírgula
                continue
            if len(nome.split()) >= 2 and INSTITUICAO_RE.search(nome):
                avisos.append(f"“{nome}” estava na lista de autores, mas parece uma instituição; foi removido.")
                continue
            autores.append({"n": nome, "m": list(pend) if prefixo else []})
    return autores, avisos


def parse_afiliacoes(v):
    """'[1] INPE; [2] USP' / '[1,5] INPE, Brasil, [2]IPGP' / 'INPE; USP' -> [{'m': '1', 't': 'INPE'}, ...]"""
    s = txt(v)
    if not s:
        return []

    def limpa(t):
        return re.sub(r"\s+", " ", t).strip(" \t;,").rstrip(".").strip()

    itens = []
    ms = list(BRACKET_RE.finditer(s))
    if not ms:
        itens = [{"m": "", "t": limpa(p)} for p in re.split(r"[;\n]+", s) if limpa(p)]
    else:
        if limpa(s[:ms[0].start()]):
            itens.append({"m": "", "t": limpa(s[:ms[0].start()])})
        for i, m in enumerate(ms):
            fim = ms[i + 1].start() if i + 1 < len(ms) else len(s)
            t = limpa(s[m.end():fim])
            if t:
                itens.append({"m": ",".join(marcas_de(m.group(1))), "t": t})
    # instituições com o mesmo texto viram uma linha só, com todos os números ("3,4,5,7 INPE")
    unicos, por_texto = [], {}
    for it in itens:
        k = re.sub(r"\W+", " ", norm_espacos(it["t"])).strip()
        if k in por_texto:
            ja = por_texto[k]
            marcas = ja["m"].split(",") if ja["m"] else []
            marcas += [x for x in (it["m"].split(",") if it["m"] else []) if x not in marcas]
            ja["m"] = ",".join(marcas)
        else:
            por_texto[k] = it
            unicos.append(it)
    return unicos


def mesma_pessoa(a, b):
    """Compara 'R. de Jesus' com 'Rodolfo de Jesus', 'Moura, R.' com 'Rafael Moura' (ignora acento/maiúscula)."""
    def toks(x):
        return [t for t in re.split(r"[^a-z0-9]+", norm_espacos(x)) if t]
    ta, tb = toks(a), toks(b)
    if not ta or not tb:
        return False
    sobrenome_comum = {t for t in ta if len(t) >= 3} & {t for t in tb if len(t) >= 3}
    ia, ib = {t[0] for t in ta}, {t[0] for t in tb}
    return bool(sobrenome_comum) and (ia <= ib or ib <= ia)


def norm_espacos(x):
    x = unicodedata.normalize("NFKD", str(x or ""))
    return "".join(c for c in x if not unicodedata.combining(c)).lower()


def paragrafos(v):
    """Cada linha do abstract vira um parágrafo (o Forms usa Enter); linha que termina no meio da frase é emendada."""
    linhas = [re.sub(r"\s+", " ", l).strip() for l in txt(v).split("\n")]
    ps = []
    for l in linhas:
        if not l:
            continue
        if ps and not re.search(r"[.!?:;)\]\"”’]$", ps[-1]) and l[0].islower():
            ps[-1] += " " + l
        else:
            ps.append(l)
    return ps


def id_de(ts, linha, aba):
    """Chave estável da apresentação: a data/hora de envio do formulário (não muda se ordenarem a planilha)."""
    if isinstance(ts, dt.datetime):
        return "T" + ts.strftime("%Y%m%d-%H%M%S") + f"{ts.microsecond // 1000:03d}"
    dig = re.sub(r"\D", "", str(ts or ""))
    return ("T" + dig) if len(dig) >= 8 else f"L{linha}" + (f"-{aba}" if aba else "")


def inteiro(v):
    m = re.fullmatch(r"\s*(\d+)(?:\.0+)?\s*", str(v)) if v not in (None, "") else None
    return int(m.group(1)) if m else None


def parse_poster(v):
    """'P07' / 'p7' / 'P 07' / '07' / 7 -> 7  (None se não for um número de pôster)."""
    m = re.fullmatch(r"\s*p?\s*[-.]?\s*0*(\d+)\s*", txt(v), re.I) if txt(v) else None
    return int(m.group(1)) if m else None


def faixas(numeros, prefixo="P", largura=2):
    """[1,2,3,5,9,10] -> 'P01–P03, P05, P09–P10'"""
    out, ns = [], sorted(numeros)
    i = 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
            j += 1
        a = f"{prefixo}{ns[i]:0{largura}d}"
        out.append(a if j == i else f"{a}–{prefixo}{ns[j]:0{largura}d}")
        i = j + 1
    return ", ".join(out)


def ler_planilha(caminho, abas=None):
    """Lê uma ou mais abas. Devolve (registros, quantas_inscricoes_sem_apresentacao_foram_ignoradas)."""
    try:
        wb = load_workbook(caminho, data_only=True)
    except Exception as e:                           # arquivo aberto, corrompido, etc.
        raise ErroPlanilha(f"Não consegui abrir {caminho}: {e}")
    if abas:
        for a in abas:
            if a not in wb.sheetnames:
                raise ErroPlanilha(f"A planilha não tem a aba '{a}'. Abas: {', '.join(wb.sheetnames)}")
        planilhas = [wb[a] for a in abas]
    else:
        planilhas = [wb["Programa"] if "Programa" in wb.sheetnames else wb.worksheets[0]]

    registros, ignoradas = [], 0
    for ws in planilhas:
        linhas = list(ws.iter_rows(values_only=True))
        if not linhas:
            raise ErroPlanilha(f"A aba '{ws.title}' está vazia.")
        colunas = {}
        for i, cel in enumerate(linhas[0]):
            k = norm(cel)
            if not k:
                continue
            for campo, apelidos in CAMPOS.items():
                if campo in colunas:
                    continue
                if any(k == a or (len(a) >= 8 and k.startswith(a)) for a in apelidos):
                    colunas[campo] = i
                    break
        faltam = [c for c in OBRIGATORIOS if c not in colunas]
        if faltam:
            raise ErroPlanilha(
                f"Aba '{ws.title}': faltam colunas no cabeçalho (linha 1): " + ", ".join(faltam)
                + ". Colunas encontradas: " + ", ".join(str(c).strip() for c in linhas[0] if c)
            )

        for n, row in enumerate(linhas[1:], start=2):
            def cel(campo):
                i = colunas.get(campo)
                return row[i] if i is not None and i < len(row) else None

            if all(v is None or str(v).strip() == "" for v in row):
                continue                             # linha vazia
            tipo_bruto = txt(cel("tipo"))
            titulo = txt(cel("titulo"))
            if not titulo and not tipo_bruto and not txt(cel("autores")) and not txt(cel("abstract")):
                ignoradas += 1                       # inscrição sem apresentação (só assiste)
                continue
            autores, avisos_aut = parse_autores(cel("autores"))
            apresentador = ajusta_caixa(re.sub(r"\s+", " ", txt(cel("apresentador"))))
            if not autores and apresentador:
                autores = [{"n": apresentador, "m": []}]
            sessao = txt(cel("sessao"))
            tipo = tipo_de(tipo_bruto)
            # número do pôster: coluna "Poster" ou, se não houver, a coluna Order
            pbruto = (txt(cel("pnum")) or txt(cel("ordem"))) if tipo == "poster" else ""
            registros.append({
                "aba": ws.title if len(planilhas) > 1 else "",
                "linha": n,
                "id": txt(cel("id")) or id_de(cel("timestamp"), n, ws.title if len(planilhas) > 1 else ""),
                "tipo_bruto": tipo_bruto,
                "tipo": tipo,
                "poster_bruto": pbruto,
                "poster_no": parse_poster(pbruto) if pbruto else None,
                "sessao": sessao,
                "dia_sessao": dia_da_sessao(sessao),
                "dia": parse_dia(cel("dia")) if txt(cel("dia")) else None,
                "dia_bruto": txt(cel("dia")),
                "ordem": inteiro(cel("ordem")),
                "ordem_bruta": txt(cel("ordem")),
                "inicio": hhmm(cel("inicio")),
                "fim": hhmm(cel("fim")),
                "titulo": re.sub(r"\s+", " ", titulo),
                "apresentador": apresentador,
                "autores": autores,
                "afiliacoes": parse_afiliacoes(cel("afiliacoes")),
                "abstract": paragrafos(cel("abstract")),
                "avisos_leitura": avisos_aut,
            })
    return registros, ignoradas


# ============================================================================
# Leitura da página (grade de horários)
# ============================================================================
def ler_grade(linhas):
    """Devolve ({(dia, inicio): info}, {dia: {'data':..., 'tema':...}})."""
    dia, grade, meta = None, {}, {}
    for i, ln in enumerate(linhas):
        if POSTER_SESSION_RE.search(ln):
            dia = None                                   # daqui para baixo são pôsteres, não horários
            continue
        m = DAY_RE.search(ln)
        if m:
            dia = int(m.group(1))
            meta.setdefault(dia, {})
            continue
        if dia is None:
            continue
        m = re.search(r'class="day-date">([^<]+)<', ln)
        if m:
            meta[dia]["data"] = html.unescape(m.group(1).strip())
        m = re.search(r'class="day-theme">([^<]+)<', ln)
        if m:
            meta[dia]["tema"] = html.unescape(m.group(1).strip())
        m = SLOT_RE.match(ln)
        if not m:
            continue
        cls = m.group("cls")
        if "break-slot" in cls or "round-slot" in cls:
            continue
        t = TEMPO_RE.match(m.group("tempo").strip())
        if not t:
            continue
        ini = hhmm(t.group(1))
        grade[(dia, ini)] = {
            "linha": i,
            "fim": hhmm(t.group(2)),
            "tipo": "invited" if "invited-slot" in cls else "oral",
            "tempo": m.group("tempo").strip(),
            "indent": m.group("indent"),
        }
    return grade, meta


def sessao_do_tema(tema, meta):
    """'Modeling' -> 1 (o dia da agenda que tem esse tema); None se nenhum dia tiver."""
    alvo = norm(tema)
    for d, info in sorted(meta.items()):
        if norm(info.get("tema", "")) == alvo:
            return d
    return None


def achar_css_posters(texto):
    """(ini, fim) em caracteres do CSS dos pôsteres: do comentário '/* POSTER SESSION' até a última regra
    .poster-* que vem em seguida. None se não achar."""
    m = re.search(r"^[ \t]*/\*\s*POSTER SESSION", texto, re.M)
    if not m:
        return None
    i, fim, n = m.start(), None, len(texto)
    while i < n:
        while i < n and texto[i].isspace():
            i += 1
        if texto.startswith("/*", i):
            j = texto.find("*/", i)
            if j < 0:
                break
            i = j + 2
            continue
        if texto.startswith("</", i):
            break
        j = texto.find("{", i)
        if j < 0:
            break
        prof, k = 0, j
        while k < n:
            if texto[k] == "{":
                prof += 1
            elif texto[k] == "}":
                prof -= 1
                if prof == 0:
                    break
            k += 1
        item = texto[i:k + 1]
        if "</style" in item or ".poster-" not in item:
            break
        fim = i = k + 1
    return (m.start(), fim) if fim else None


def periodo_do_evento(meta):
    """'19–21 October' a partir das datas dos dias da agenda ('' se não der para montar)."""
    ms = [re.match(r"(\d+)\s+(\S+)", meta[d].get("data", "")) for d in sorted(meta)]
    ms = [m for m in ms if m]
    if not ms:
        return ""
    a, b = ms[0], ms[-1]
    if a.group(2) == b.group(2):
        return f"{a.group(1)}\u2013{b.group(1)} {a.group(2)}" if a.group(1) != b.group(1) else f"{a.group(1)} {a.group(2)}"
    return f"{a.group(1)} {a.group(2)} \u2013 {b.group(1)} {b.group(2)}"


def ler_posters(linhas, meta):
    """
    Lê a seção "Poster session" da página. Devolve (blocos_p, secao).
      blocos_p -> um bloco por tema da AGENDA (Modeling = sessão 1, Networks = sessão 2...), na ordem dos dias:
                  [{'tema', 'sessao', 'ocupados', 'reservados'}]. 'ocupados' = pôsteres que a página já tinha
                  desse tema; 'reservados' = slots TBD numerados que o usuário deixou nele (só no formato novo).
      secao    -> o miolo de <div class="poster-days"> que o script reescreve por inteiro:
                  {'ini', 'fim', 'indent', 'antigo', 'cab', 'temas'}; 'antigo' = a página ainda está no formato de
                  3 colunas (Day 01/02/03); 'cab' = linhas do título + parágrafo; 'temas' = temas que a página tinha.
                  None se a página não tem a seção de pôsteres.
    """
    ini_sec = next((i for i, ln in enumerate(linhas) if POSTER_SESSION_RE.search(ln)), None)
    if ini_sec is None:
        return [], None
    c_ini = next((i for i in range(ini_sec, len(linhas)) if POSTER_DAYS_RE.search(linhas[i])), None)
    if c_ini is None:
        return [], None
    prof, c_fim = 0, None
    for i in range(c_ini, len(linhas)):
        prof += linhas[i].count("<div") - linhas[i].count("</div>")
        if prof <= 0:
            c_fim = i
            break
    if c_fim is None:
        return [], None

    h3 = next((i for i in range(ini_sec, c_ini) if "<h3" in linhas[i]), None)
    p_fim = next((i for i in range(h3, c_ini) if "</p>" in linhas[i]), None) if h3 is not None else None
    cab = (h3, p_fim) if p_fim is not None else None

    miolo = linhas[c_ini + 1:c_fim]
    antigo = any(re.search(r"poster-day-top|<!--\s*POSTERS\s+DAY", ln) for ln in miolo)
    existentes, temas, tema, i = {}, [], None, 0
    while i < len(miolo):
        m = POSTER_TEMA_RE.search(miolo[i])
        if m:
            tema = html.unescape(m.group(1).strip())
        elif POSTER_LISTA_RE.search(miolo[i]) and tema is not None:
            fim = i + 1
            while fim < len(miolo) and not FECHA_DIV_RE.match(miolo[fim]):
                fim += 1
            oc = rs = 0
            for k in range(i + 1, min(fim, len(miolo))):
                if POSTER_CARD_RE.match(miolo[k]):
                    if SLOT_RESERVADO_RE.search(miolo[k]):
                        rs += 1                          # slot numerado em TBD: vaga que o usuário deixou
                    elif "has-talk" in miolo[k]:
                        oc += 1                          # pôster que a geração anterior colocou
            temas.append(tema)
            existentes[norm(tema)] = (oc, rs)
            tema, i = None, fim
        i += 1

    blocos = []
    for d, info in sorted(meta.items()):
        t = info.get("tema", "")
        if not t:
            continue
        oc, rs = (0, 0) if antigo else existentes.get(norm(t), (0, 0))   # TBD do formato antigo não valem como reserva
        blocos.append({"tema": t, "sessao": d, "ocupados": oc, "reservados": rs})
    secao = {"ini": c_ini + 1, "fim": c_fim, "indent": re.match(r"[ \t]*", linhas[c_ini]).group(),
             "antigo": antigo, "cab": cab, "temas": temas}
    return blocos, secao


# ============================================================================
# Validação
# ============================================================================
def rotulo(r):
    quem = r["apresentador"] or (r["autores"][0]["n"] if r["autores"] else "")
    return (f"linha {r['linha']}" + (f" da aba {r['aba']}" if r["aba"] else "") + (f" ({quem})" if quem else ""))


def validar(registros, grade, ordem_das_linhas=False, estrito=False, blocos_p=None, meta=None):
    """
    Devolve (erros, avisos, colocados, sem_horario, excluidos, colocados_p).
      erros       -> só no modo estrito: impedem de escrever a página
      excluidos   -> linhas com problema: ficam FORA da agenda (a página é escrita com o que está certo)
      avisos      -> o que vale conferir na planilha
      colocados   -> {(dia, inicio): registro}
      sem_horario -> orais/convidados que ainda não têm horário (não é erro: aparecem quando forem escalados)
      colocados_p -> {número do pôster: registro}: pôsteres numerados em sequência (P01, P02...) na ordem dos
                     blocos da página; cada bloco guarda os seus em bloco["cartoes"] = [(número, registro)]
    """
    blocos_p, meta = blocos_p or [], meta or {}
    por_sessao = {b["sessao"]: b for b in blocos_p if b["sessao"]}
    for b in blocos_p:
        b["registros"], b["cartoes"], b["vagas_n"] = [], [], []
    erros, excluidos, avisos, colocados, ocupado, ids = [], [], [], {}, {}, {}
    for b in blocos_p:
        if b["sessao"] is None:
            avisos.append(f"o bloco de pôsteres “{b['tema']}” da página não tem o mesmo nome do tema de nenhum dia "
                          "da agenda; nenhum pôster entra nele.")

    def falha(msg):
        (erros if estrito else excluidos).append(msg)

    slots, slots_dia = {}, {}
    for (d, ini), g in sorted(grade.items()):
        slots.setdefault((d, g["tipo"]), []).append(ini)
        slots_dia.setdefault(d, []).append((ini, g["tipo"]))
    pendentes = []
    titulos, apresentadores = {}, {}

    def ocupar(r, rot, dia, ini):
        g = grade.get((dia, ini))
        if g is None:
            validos = ", ".join(sorted(h for (d, h) in grade if d == dia))
            falha(f"{rot}: não existe horário {ini} no dia {dia}. Horários do dia {dia}: {validos}.")
            return
        if g["tipo"] != r["tipo"]:
            falha(f"{rot}: o horário {ini} do dia {dia} é de {g['tipo']}, mas a linha é {r['tipo']}.")
            return
        if r["fim"] and r["fim"] != g["fim"]:
            avisos.append(f"{rot}: fim {r['fim']} difere da agenda ({g['fim']}); vale o horário da agenda.")
        if (dia, ini) in ocupado:
            o = ocupado[(dia, ini)]
            falha(f"{rot}: o horário dia {dia} {ini} já está ocupado por {rotulo(o)}.")
            return
        ocupado[(dia, ini)] = r
        colocados[(dia, ini)] = r

    for r in registros:
        rot = rotulo(r)
        if r["id"] in ids:
            falha(f"{rot}: id repetido (já usado em {ids[r['id']]}).")
            continue
        ids[r["id"]] = rot

        if r["tipo"] is None:
            msg = f"tipo inválido '{r['tipo_bruto']}'" if r["tipo_bruto"] else "falta o tipo"
            falha(f"{rot}: {msg} (use Oral, Poster ou Invited).")
            continue
        if not r["titulo"]:
            falha(f"{rot}: falta o título.")
            continue
        if not r["autores"]:
            falha(f"{rot}: falta o autor.")
            continue
        if not r["abstract"]:
            avisos.append(f"{rot}: sem abstract (o título não ficará clicável).")

        # ---- qualidade dos dados (só avisos)
        if "lorem ipsum" in " ".join(r["abstract"]).lower():
            avisos.append(f"{rot}: o abstract parece texto de teste (lorem ipsum).")
        if len(r["titulo"]) > 3 and r["titulo"] == r["titulo"].upper():
            avisos.append(f"{rot}: título todo em maiúsculas: “{r['titulo']}”.")
        for a in r["avisos_leitura"]:
            avisos.append(f"{rot}: {a}")
        if r["apresentador"] and r["autores"] and not mesma_pessoa(r["apresentador"], r["autores"][0]["n"]):
            avisos.append(f"{rot}: o apresentador “{r['apresentador']}” não é o primeiro autor "
                          f"(“{r['autores'][0]['n']}”); a janela do abstract vai mostrar “Presenting author”.")
        chave_t = norm(r["titulo"])
        if chave_t and chave_t in titulos:
            avisos.append(f"{rot}: título igual ao de {titulos[chave_t]}.")
        titulos.setdefault(chave_t, rot)
        if r["apresentador"]:
            apresentadores.setdefault(norm_espacos(r["apresentador"]).strip(), []).append(r)

        if r["tipo"] == "poster":
            if not blocos_p:
                continue                                 # a página não tem a seção de pôsteres
            if r["poster_bruto"] and r["poster_no"] is None:
                falha(f"{rot}: número de pôster inválido '{r['poster_bruto']}' (use 1, 2, 3... ou P01, P02...; "
                      "ele só define a ordem dentro do tema).")
                continue
            sess = r["dia_sessao"] or r["dia"]           # pôster não tem horário: vale o tema (ou o Day, se faltar)
            if sess is None:
                falha(f"{rot}: falta a sessão temática do pôster; não sei em qual bloco da página ele entra.")
                continue
            bloco = por_sessao.get(sess)
            if bloco is None:
                tema = meta.get(sess, {}).get("tema", "")
                avisos.append(f"{rot}: a seção de pôsteres da página não tem bloco para a sessão {sess}"
                              + (f" ({tema})" if tema else "")
                              + "; o pôster ficou de fora. Blocos da página: "
                              + ", ".join(b["tema"] for b in blocos_p) + ".")
                continue
            bloco["registros"].append(r)
            continue

        # ---- onde entra na agenda
        if r["dia_bruto"] and r["dia"] is None:
            falha(f"{rot}: dia inválido '{r['dia_bruto']}' (use 1, 2 ou 3).")
            continue
        if r["ordem_bruta"] and r["ordem"] is None:
            falha(f"{rot}: ordem inválida '{r['ordem_bruta']}' (use 1, 2, 3...).")
            continue
        dia = r["dia"] or r["dia_sessao"]            # o dia da planilha vale mais que o da sessão temática
        if not r["inicio"] and not r["ordem"]:
            pendentes.append((r, rot, dia))
            continue
        if dia is None:
            falha(f"{rot}: falta o dia (e a sessão temática não indica um).")
            continue
        if r["inicio"]:
            ini = r["inicio"]
        else:
            # Order = posição entre TODOS os horários do dia, em ordem de relógio (convidados e orais juntos)
            lista_d = slots_dia.get(dia, [])
            if r["ordem"] < 1:
                falha(f"{rot}: ordem {r['ordem']} inválida (use 1, 2, 3...).")
                continue
            if r["ordem"] > len(lista_d):
                # mais inscritos do que vagas: quem passa do limite fica de fora da agenda, sem bloquear o resto
                avisos.append(f"{rot}: ordem {r['ordem']} passa dos {len(lista_d)} horários do dia {dia}; "
                              "ficou de fora da agenda (mude o dia, passe para Poster ou ignore por enquanto).")
                continue
            ini, tipo_slot = lista_d[r["ordem"] - 1]
            if tipo_slot != r["tipo"]:
                falha(f"{rot}: no dia {dia} a ordem {r['ordem']} é o horário {ini}, que é de {tipo_slot}, "
                      f"mas a linha é {r['tipo']}.")
                continue
        ocupar(r, rot, dia, ini)

    sem_horario = []
    for r, rot, dia in pendentes:                    # 2ª passada: quem ainda não tem horário
        if ordem_das_linhas and dia is not None:
            livres = [h for h in slots.get((dia, r["tipo"]), []) if (dia, h) not in ocupado]
            if livres:
                r["provisorio"] = True
                ocupar(r, rot, dia, livres[0])
                continue
        sem_horario.append(r)

    colocados_p, n = {}, 0
    for b in blocos_p:                               # números em sequência contínua, na ordem dos blocos da página
        # a coluna Poster/Order só decide a ordem dentro do tema (sem ela, vale a ordem da planilha: sort estável)
        ordenados = sorted(b["registros"], key=lambda x: (x["poster_no"] is None, x["poster_no"] or 0))
        usados = {}
        for r in ordenados:
            n += 1
            r["poster_n"], r["poster_tema"] = n, b["tema"]
            colocados_p[n] = r
            b["cartoes"].append((n, r))
            if r["poster_no"] is not None:
                if r["poster_no"] in usados:
                    avisos.append(f"o número de pôster {r['poster_no']} aparece em {rotulo(usados[r['poster_no']])} "
                                  f"e em {rotulo(r)}; a ordem entre eles segue a planilha.")
                usados.setdefault(r["poster_no"], r)
        # Slots reservados (TBD) deixados na página: ficam depois dos pôsteres do bloco e são ocupados, um a um,
        # pelos pôsteres NOVOS da planilha (os que passam do que a página já tinha preenchido).
        novos = max(0, len(b["registros"]) - b["ocupados"])
        for _ in range(max(0, b["reservados"] - novos)):
            n += 1
            b["vagas_n"].append(n)

    for lista_r in apresentadores.values():
        if len(lista_r) > 1:
            avisos.append("o mesmo apresentador aparece em " + str(len(lista_r)) + " trabalhos: "
                          + "; ".join(rotulo(x) for x in lista_r) + ".")
    return erros, avisos, colocados, sem_horario, excluidos, colocados_p


# ============================================================================
# Geração do HTML
# ============================================================================
def esc(s):
    return html.escape(s, quote=True)


def autor_exibicao(r):
    autores = r["autores"]
    primeiro = esc(autores[0]["n"])
    return primeiro + (" <em>et al.</em>" if len(autores) > 1 else "")


def render_slot(g, r, original=False):
    """original=True: página sem nenhuma apresentação -> devolve exatamente o HTML original do site."""
    invited = g["tipo"] == "invited"
    base = "slot" + (" invited-slot" if invited else "")
    badge = (
        '<span class="type-badge invited">Invited</span>' if invited
        else '<span class="type-badge talk">Talk</span>'
    )
    ini = f'{g["indent"]}<div class="{base}'
    time = f'<span class="slot-time">{g["tempo"]}</span>'
    if original:
        return f'{ini}">{time}<span class="slot-content">{badge}<br>TBD</span></div>'
    meta = f'<span class="slot-meta">{time}{badge}</span>'      # horário + selo, na coluna da esquerda
    if r is None:
        return f'{ini}">{meta}<span class="slot-content tbd">TBD</span></div>'
    miolo = (
        f'<span class="talk-title">{esc(r["titulo"])}</span>'
        f'<span class="talk-author">{autor_exibicao(r)}</span>'
    )
    if r["abstract"]:
        alvo = (
            f'<button type="button" class="talk-link" data-id="{esc(r["id"])}" '
            f'title="{esc(r["titulo"])} — click to read the abstract" aria-haspopup="dialog">{miolo}</button>'
        )
    else:
        alvo = f'<span class="talk-link">{miolo}</span>'
    return f'{ini} has-talk">{meta}<span class="slot-content">{alvo}</span></div>'


def render_poster_vazio(indent):
    """Bloco de tema sem nenhum pôster na planilha: uma linha avisando (a caixa nunca fica vazia)."""
    return (f'{indent}<div class="poster-card"><span class="poster-no">&ndash;</span>'
            '<span class="poster-content tbd">To be announced</span></div>')


def render_poster_livre(indent, numero):
    """Slot reservado: cartão numerado, ainda sem pôster (cinza, "TBD")."""
    rot = f"P{numero:02d}"
    return (f'{indent}<div class="poster-card" data-poster="{rot}"><span class="poster-no">{rot}</span>'
            '<span class="poster-content tbd">TBD</span></div>')


def render_poster(indent, numero, r):
    """Uma linha <div class="poster-card"...> da seção Poster session (um cartão por pôster)."""
    rot = f"P{numero:02d}"
    ini = f'{indent}<div class="poster-card'
    no = f'<span class="poster-no">{rot}</span>'
    miolo = (
        f'<span class="talk-title">{esc(r["titulo"])}</span>'
        f'<span class="talk-author">{autor_exibicao(r)}</span>'
    )
    if r["abstract"]:
        alvo = (
            f'<button type="button" class="talk-link" data-id="{esc(r["id"])}" '
            f'title="{esc(r["titulo"])} — click to read the abstract" aria-haspopup="dialog">{miolo}</button>'
        )
    else:
        alvo = f'<span class="talk-link">{miolo}</span>'
    return f'{ini} has-talk" data-poster="{rot}">{no}<span class="poster-content">{alvo}</span></div>'


def render_blocos_poster(blocos_p, secao):
    """
    O miolo de <div class="poster-days">: um bloco por tema que tem pôster (ou slot reservado), na ordem dos dias
    da agenda. Tema sem nenhum pôster não aparece. Se nenhum tema tem pôster, ficam os temas que a página já tinha
    (ou todos, se não tinha nenhum), cada um com "To be announced".
    """
    base = secao["indent"]
    com_conteudo = [b for b in blocos_p if b.get("cartoes") or b.get("vagas_n")]
    if com_conteudo:
        usar = com_conteudo
    else:
        tinha = {norm(t) for t in secao["temas"]}
        usar = [b for b in blocos_p if norm(b["tema"]) in tinha] or blocos_p
    saida, ci = [], base + "    "
    for b in usar:
        cartoes = ([render_poster(ci, n, r) for n, r in b.get("cartoes", [])]
                   + [render_poster_livre(ci, n) for n in b.get("vagas_n", [])])
        saida += ([f"{base}<!-- POSTERS: {esc(b['tema']).upper()} -->",
                   f'{base}<div class="poster-day">',
                   f'{base}  <div class="poster-day-head">',
                   f'{base}    <div class="poster-day-theme">{esc(b["tema"])}</div>',
                   f'{base}  </div>',
                   f'{base}  <div class="poster-list">']
                  + (cartoes or [render_poster_vazio(ci)])
                  + [f'{base}  </div>', f'{base}</div>', ""])
    return saida


def cabecalho_posters(meta, indent):
    """Título + parágrafo da seção, no texto do formato novo (só usado ao migrar uma página no formato antigo)."""
    periodo = periodo_do_evento(meta)
    return [f'{indent}<h3>Poster <em>session</em> <span class="poster-sub">(remain on display every day)</span></h3>',
            f"{indent}<p>",
            f"{indent}  Posters are organized by scientific theme and remain on display every day throughout the "
            f"workshop" + (f" ({periodo})" if periodo else "")
            + ", so you can visit any of them at any time. Click a poster title to read its abstract.",
            f"{indent}</p>"]


def json_seguro(dados):
    """Um id por linha; escapa < > & para poder ficar dentro de <script>."""
    partes = []
    for k, v in dados.items():
        s = json.dumps(v, ensure_ascii=False, separators=(",", ":"))
        s = (s.replace("&", "\\u0026").replace("<", "\\u003c").replace(">", "\\u003e")
              .replace("\u2028", "\\u2028").replace("\u2029", "\\u2029"))
        partes.append(f"{json.dumps(k)}:{s}")
    return "{\n" + ",\n".join(partes) + "\n}"


def dados_modal(colocados, meta, grade, colocados_p=None):
    rotulo_tipo = {"invited": "Invited talk", "oral": "Oral presentation"}
    dados = {}
    for (dia, ini), r in sorted(colocados.items()):
        if not r["abstract"]:
            continue
        g = meta.get(dia, {})
        # só mostra "Presenting author" quando o apresentador não é o primeiro autor
        mostrar_apres = r["apresentador"] and not mesma_pessoa(r["apresentador"], r["autores"][0]["n"])
        dados[r["id"]] = {
            "tipo": r["tipo"],
            "rotulo": rotulo_tipo[r["tipo"]],
            "dia": dia,
            "data": g.get("data", ""),
            "tema": g.get("tema", ""),
            "inicio": ini,
            "fim": grade[(dia, ini)]["fim"],
            "titulo": r["titulo"],
            "autores": [{"n": a["n"], "m": ",".join(a["m"])} for a in r["autores"]],
            "afiliacoes": r["afiliacoes"],
            "apresentador": r["apresentador"] if mostrar_apres else "",
            "abstract": r["abstract"],
        }
    for n, r in sorted((colocados_p or {}).items()):         # pôsteres: sem dia nem horário (expostos todos os dias)
        if not r["abstract"]:
            continue
        mostrar_apres = r["apresentador"] and not mesma_pessoa(r["apresentador"], r["autores"][0]["n"])
        dados[r["id"]] = {
            "tipo": "poster",
            "rotulo": f"Poster P{n:02d}",
            "dia": "",
            "data": "",
            "tema": r.get("poster_tema", ""),
            "inicio": "",
            "fim": "",
            "titulo": r["titulo"],
            "autores": [{"n": a["n"], "m": ",".join(a["m"])} for a in r["autores"]],
            "afiliacoes": r["afiliacoes"],
            "apresentador": r["apresentador"] if mostrar_apres else "",
            "abstract": r["abstract"],
        }
    return dados


CSS = r"""
  /* ── Agenda com apresentações ── */
  html.abs-open { overflow: hidden; }

  /* Coluna da esquerda: horário + selo (Talk / Invited / Coffee break / Round table), empilhados
     e centralizados na vertical dentro do cartão. */
  .slot { align-items: center; }
  .slot:not(.break-slot):not(.round-slot) { padding-top: 11px; padding-bottom: 11px; }
  .slot-meta { display: flex; flex-direction: column; align-items: flex-start; gap: 6px; }
  .slot-meta .slot-time { padding: 0; }
  .slot-meta .type-badge {
    display: inline-block; font-family: 'DM Mono', monospace; font-size: 8.5px; line-height: 1.5;
    letter-spacing: 0.14em; text-transform: uppercase; padding: 1px 6px; border-radius: 2px;
  }
  .slot.break-slot, .slot.round-slot {
    display: flex; flex-direction: column; align-items: flex-start; justify-content: center; gap: 6px;
  }
  .slot.break-slot .slot-time, .slot.round-slot .slot-time { padding: 0; }
  .slot.round-slot .slot-content .type-badge { margin-bottom: 0; }
  .slot-content.tbd { color: var(--muted); }

  .talk-link {
    display: block; width: 100%; margin: 0; padding: 0; border: 0;
    background: none; font: inherit; color: inherit; text-align: left; cursor: default;
  }
  button.talk-link { cursor: pointer; }
  .talk-title {
    display: block; font-size: 12.5px; font-weight: 500; line-height: 1.38;
    color: #f0f6ff; transition: color 0.15s;
  }
  .talk-author { display: block; margin-top: 3px; font-size: 11.5px; line-height: 1.4; color: var(--muted); }
  .talk-author em { font-style: italic; }
  /* o "+" fica na linha do autor, que nunca é cortada (o título pode ser) */
  button.talk-link .talk-author::after {
    content: "\00a0\00a0+"; font-family: 'DM Mono', monospace; font-size: 11px; font-style: normal;
    color: var(--electric-dim); transition: color 0.15s;
  }
  button.talk-link:hover .talk-title,
  button.talk-link:focus-visible .talk-title,
  button.talk-link:hover .talk-author::after,
  button.talk-link:focus-visible .talk-author::after { color: var(--aurora); }
  button.talk-link:focus-visible { outline: 2px solid var(--electric); outline-offset: 4px; border-radius: 2px; }

  /* Altura fixa (3 colunas): todos os cartões de apresentação (oral e convidado) têm o mesmo tamanho,
     então a mesma linha fica alinhada nos três dias. Para ajustar tudo de uma vez, mude --slot-h.
     Título maior que o cartão é cortado com "…" (o título completo aparece ao clicar).
     A coluna do horário é mais estreita aqui para sobrar mais largura para o título. */
  @media (min-width: 861px) {
    .agenda-days { --slot-h: 152px; }
    .slot:not(.break-slot):not(.round-slot) {
      box-sizing: border-box; height: var(--slot-h); overflow: hidden;
      grid-template-columns: 78px 1fr; column-gap: 8px; padding-right: 12px;
    }
    .slot .talk-title {
      display: -webkit-box; -webkit-box-orient: vertical; -webkit-line-clamp: 6; line-clamp: 6; overflow: hidden;
    }
  }
  /* telas largas só um pouco: colunas mais estreitas, título quebra em mais linhas */
  @media (min-width: 861px) and (max-width: 999px) {
    .agenda-days { --slot-h: 184px; }
    .slot .talk-title { -webkit-line-clamp: 7; line-clamp: 7; }
  }

  .demo-banner {
    margin: 22px 0 0; padding: 10px 16px; border: 1px dashed rgba(255, 107, 107, 0.6);
    background: rgba(255, 107, 107, 0.07); color: #ff9c9c; border-radius: 3px;
    font-family: 'DM Mono', monospace; font-size: 11px; letter-spacing: 0.12em; text-transform: uppercase;
  }

  /* Janela do abstract */
  .abs-dialog {
    margin: auto; padding: 0; width: min(720px, calc(100vw - 24px)); max-width: none;
    max-height: min(88vh, 780px); overflow: hidden; color: var(--ink);
    background: var(--navy-mid); border: 1px solid var(--border-strong);
    border-top: 3px solid var(--electric); border-radius: 6px;
    box-shadow: 0 28px 90px rgba(0, 0, 0, 0.6);
  }
  .abs-dialog[data-tipo="invited"] { border-top-color: var(--gold); }
  .abs-dialog::backdrop { background: rgba(4, 10, 24, 0.74); backdrop-filter: blur(3px); }
  .abs-card { position: relative; max-height: min(88vh, 780px); overflow-y: auto; padding: 34px 40px 36px; }
  .abs-close {
    position: absolute; top: 14px; right: 14px; width: 34px; height: 34px; padding: 0;
    display: flex; align-items: center; justify-content: center; cursor: pointer;
    color: var(--ink-soft); background: rgba(79, 195, 247, 0.06);
    border: 1px solid var(--border-strong); border-radius: 50%;
    transition: color 0.2s, background 0.2s, border-color 0.2s;
  }
  .abs-close:hover, .abs-close:focus-visible { color: var(--navy); background: var(--aurora); border-color: var(--aurora); }
  .abs-eyebrow {
    margin: 0 44px 14px 0; font-family: 'DM Mono', monospace; font-size: 10px;
    letter-spacing: 0.16em; line-height: 1.7; text-transform: uppercase; color: var(--electric);
  }
  .abs-dialog[data-tipo="invited"] .abs-eyebrow { color: var(--gold); }
  .abs-dialog[data-tipo="poster"] { border-top-color: var(--aurora); }
  .abs-dialog[data-tipo="poster"] .abs-eyebrow { color: var(--aurora); }
  .abs-title {
    margin: 0 0 16px; font-family: 'Cormorant Garamond', serif; font-size: 30px;
    font-weight: 600; line-height: 1.2; color: var(--white);
  }
  .abs-authors { font-size: 15px; line-height: 1.6; color: #d8e8f8; }
  .abs-affils { margin-top: 6px; font-size: 12.5px; line-height: 1.6; color: var(--muted); }
  .abs-authors sup, .abs-affils sup {
    font-family: 'DM Mono', monospace; font-size: 0.68em; line-height: 0; color: var(--electric); margin-left: 1px;
  }
  .abs-affils sup { margin: 0 3px 0 0; }
  .abs-presenter {
    margin-top: 10px; font-family: 'DM Mono', monospace; font-size: 10.5px;
    letter-spacing: 0.08em; text-transform: uppercase; color: var(--aurora);
  }
  .abs-label {
    display: flex; align-items: center; gap: 12px; margin: 26px 0 14px;
    font-family: 'DM Mono', monospace; font-size: 10px; letter-spacing: 0.25em;
    text-transform: uppercase; color: var(--electric-dim);
  }
  .abs-label::after { content: ''; flex: 1; height: 1px; background: var(--border); }
  .abs-body p { margin: 0 0 14px; font-size: 15px; line-height: 1.8; color: #c0cfe8; }
  .abs-body p:last-child { margin-bottom: 0; }
  @media (max-width: 560px) {
    .abs-card { padding: 28px 22px; }
    .abs-title { font-size: 24px; }
  }
"""

MODAL_HTML = r"""<dialog id="abstract-dialog" class="abs-dialog" aria-labelledby="abs-title">
  <div class="abs-card">
    <button type="button" class="abs-close" aria-label="Close">
      <svg width="14" height="14" viewBox="0 0 24 24" fill="none" xmlns="http://www.w3.org/2000/svg"><path d="M18 6L6 18M6 6l12 12" stroke="currentColor" stroke-width="2" stroke-linecap="round"/></svg>
    </button>
    <div class="abs-eyebrow"></div>
    <h3 class="abs-title" id="abs-title"></h3>
    <div class="abs-authors"></div>
    <div class="abs-affils"></div>
    <div class="abs-presenter"></div>
    <div class="abs-label">Abstract</div>
    <div class="abs-body"></div>
  </div>
</dialog>
<script type="application/json" id="programa-data">
__DADOS__
</script>
<script>
(function () {
  var dlg = document.getElementById('abstract-dialog');
  var store = document.getElementById('programa-data');
  if (!dlg || !store) return;
  var data = {};
  try { data = JSON.parse(store.textContent); } catch (e) { return; }
  var trigger = null;
  function q(sel) { return dlg.querySelector(sel); }

  function fill(id) {
    var t = data[id];
    if (!t) return false;
    dlg.setAttribute('data-tipo', t.tipo);
    q('.abs-eyebrow').textContent = [t.dia ? 'Day ' + t.dia : '', t.tema, t.data, t.inicio ? t.inicio + '\u2013' + t.fim : '', t.rotulo]
      .filter(Boolean).join(' \u00b7 ');
    q('.abs-title').textContent = t.titulo;
    var au = q('.abs-authors');
    au.textContent = '';
    (t.autores || []).forEach(function (a, i) {
      if (i) au.appendChild(document.createTextNode('; '));
      au.appendChild(document.createTextNode(a.n));
      if (a.m) { var sp = document.createElement('sup'); sp.textContent = a.m; au.appendChild(sp); }
    });
    var aff = q('.abs-affils');
    aff.textContent = '';
    (t.afiliacoes || []).forEach(function (a) {
      var d = document.createElement('div');
      if (a.m) { var sp = document.createElement('sup'); sp.textContent = a.m; d.appendChild(sp); }
      d.appendChild(document.createTextNode((a.m ? ' ' : '') + a.t));
      aff.appendChild(d);
    });
    var pr = q('.abs-presenter');
    pr.textContent = t.apresentador ? 'Presenting author: ' + t.apresentador : '';
    pr.style.display = t.apresentador ? '' : 'none';
    var body = q('.abs-body');
    body.textContent = '';
    (t.abstract || []).forEach(function (p) {
      var el = document.createElement('p');
      el.textContent = p;
      body.appendChild(el);
    });
    q('.abs-card').scrollTop = 0;
    return true;
  }

  function onClosed() {
    document.documentElement.classList.remove('abs-open');
    if (trigger && trigger.focus) trigger.focus();
    trigger = null;
  }

  function openDialog(id, el) {
    if (!fill(id)) return;
    trigger = el || null;
    document.documentElement.classList.add('abs-open');
    if (typeof dlg.showModal === 'function') { if (!dlg.open) dlg.showModal(); }
    else { dlg.setAttribute('open', ''); }
  }

  function closeDialog() {
    if (typeof dlg.close === 'function') { dlg.close(); }
    else { dlg.removeAttribute('open'); onClosed(); }
  }

  dlg.addEventListener('close', onClosed);
  document.addEventListener('click', function (e) {
    var b = e.target.closest ? e.target.closest('.talk-link[data-id]') : null;
    if (b) { openDialog(b.getAttribute('data-id'), b); return; }
    var aberto = dlg.open || dlg.hasAttribute('open');
    if (aberto && (e.target === dlg || (e.target.closest && e.target.closest('.abs-close')))) closeDialog();
  });
})();
</script>"""


def substituir_bloco(texto, ini, fim, novo, ancora):
    """Troca (ou insere, ou remove) o bloco delimitado pelos marcadores ini/fim."""
    padrao = re.compile(r"[ \t]*" + re.escape(ini) + r".*?" + re.escape(fim) + r"[ \t]*\n*", re.S)
    if padrao.search(texto):
        return padrao.sub(lambda m: novo or "", texto, count=1)
    if not novo:
        return texto
    m = ancora.search(texto)
    if not m:
        raise ErroPlanilha(f"Não achei onde inserir o bloco '{ini}' na página.")
    return texto[:m.start()] + novo + texto[m.start():]


def ajustar_intro(texto, preenchida):
    """Com apresentações: frase nova entre marcadores. Sem nenhuma: volta ao texto original (sem marcadores)."""
    ini, fim = "<!-- PROGRAMA:INTRO:INICIO -->", "<!-- PROGRAMA:INTRO:FIM -->"
    padrao = re.compile(re.escape(ini) + r".*?" + re.escape(fim), re.S)
    if not preenchida:
        return padrao.sub(lambda m: INTRO_ORIGINAL, texto, count=1), True
    novo = ini + INTRO_PREENCHIDA + fim
    if padrao.search(texto):
        return padrao.sub(lambda m: novo, texto, count=1), True
    if INTRO_ORIGINAL in texto:
        return texto.replace(INTRO_ORIGINAL, novo, 1), True
    return texto, False


def montar_pagina(texto, linhas, grade, meta, colocados, aviso_html=None, colocados_p=None, blocos_p=None,
                  secao=None):
    colocados_p, blocos_p = colocados_p or {}, blocos_p or []
    tem = bool(colocados) or bool(colocados_p)
    for chave, g in grade.items():
        linhas[g["linha"]] = render_slot(g, colocados.get(chave), original=not tem)
    # pôsteres: o miolo de <div class="poster-days"> é reescrito inteiro (um bloco por tema, um cartão por
    # pôster), qualquer que seja o formato em que a página estava. Primeiro o trecho de baixo, depois o título
    # (que está acima), para os índices das linhas não se deslocarem.
    notas_p = []
    if secao:
        linhas[secao["ini"]:secao["fim"]] = render_blocos_poster(blocos_p, secao)
        if secao["antigo"] and secao["cab"]:
            a, z = secao["cab"]
            linhas[a:z + 1] = cabecalho_posters(meta, re.match(r"[ \t]*", linhas[a]).group())
    texto = "\n".join(linhas)
    if secao and secao["antigo"]:
        r = achar_css_posters(texto)
        if r:
            texto = texto[:r[0]] + POSTER_CSS_NOVO + texto[r[1]:]
            notas_p.append("A seção de pôsteres estava no formato antigo (3 colunas, Day 01/02/03): foi migrada para "
                           "o novo (um pôster por linha, um bloco por tema, sem cartões TBD antigos).")
        else:
            notas_p.append("ATENÇÃO: a seção de pôsteres estava no formato antigo, mas não achei o CSS antigo dela "
                           "(/* POSTER SESSION ... */); os blocos foram refeitos, mas o visual pode continuar em 3 colunas.")

    css = ("  /* PROGRAMA:CSS:INICIO — gerado por gerar_agenda.py; não edite à mão */"
           + CSS + "  /* PROGRAMA:CSS:FIM */\n") if tem else None
    modal = ("<!-- PROGRAMA:MODAL:INICIO — gerado por gerar_agenda.py; não edite à mão -->\n"
             + MODAL_HTML.replace("__DADOS__", json_seguro(dados_modal(colocados, meta, grade, colocados_p)))
             + "\n<!-- PROGRAMA:MODAL:FIM -->\n\n") if tem else None
    aviso = ("    <!-- PROGRAMA:AVISO:INICIO -->" + aviso_html + "<!-- PROGRAMA:AVISO:FIM -->\n\n"
             ) if (aviso_html and tem) else None

    texto = substituir_bloco(texto, "/* PROGRAMA:CSS:INICIO", "PROGRAMA:CSS:FIM */", css,
                             re.compile(r"^[ \t]*</style>", re.M))
    texto = substituir_bloco(texto, "<!-- PROGRAMA:MODAL:INICIO", "<!-- PROGRAMA:MODAL:FIM -->", modal,
                             re.compile(r"^[ \t]*</body>", re.M))
    texto = substituir_bloco(texto, "<!-- PROGRAMA:AVISO:INICIO", "<!-- PROGRAMA:AVISO:FIM -->", aviso,
                             re.compile(r'^[ \t]*<div class="agenda-days">', re.M))
    texto, ok = ajustar_intro(texto, tem)
    return texto, ok, notas_p


# ============================================================================
# Programa principal
# ============================================================================
def mapa_de_ordens(grade):
    """'1=09:00 invited · 2=09:40 oral · ...' (a grade é a mesma em todos os dias)."""
    dia = min(d for d, _ in grade)
    return " · ".join(f"{i}={h} {t}" for i, (h, t) in
                      enumerate(sorted((h, g["tipo"]) for (d, h), g in grade.items() if d == dia), start=1))


def mapa_de_posteres(blocos_p):
    """'Modeling = sessão 1 · Networks = sessão 2' (os blocos de pôsteres que existem na página)."""
    return " · ".join(f"{b['tema']} = sessão {b['sessao']}" if b["sessao"] else b["tema"] for b in blocos_p)


def main():
    ap = argparse.ArgumentParser(description="Preenche a agenda do site a partir da planilha do programa.")
    ap.add_argument("planilha", help="arquivo .xlsx (a exportação do formulário serve como está)")
    ap.add_argument("--pagina", default="index.html", help="página de entrada (padrão: index.html)")
    ap.add_argument("--saida", default=None,
                    help="arquivo de saída (padrão: sobrescreve a página, guardando cópia .bak)")
    ap.add_argument("--aba", default=None,
                    help="aba a ler; para várias, separe por vírgula (padrão: 'Programa' ou a primeira aba)")
    ap.add_argument("--ordem-das-linhas", action="store_true",
                    help="PRÉVIA: quem não tem horário/ordem entra na ordem em que aparece na planilha, "
                         "nos horários livres do dia da sua sessão temática")
    ap.add_argument("--estrito", action="store_true",
                    help="se houver qualquer problema na planilha, não escreve a página "
                         "(padrão: escreve o que está certo e lista o que ficou de fora)")
    ap.add_argument("--exemplo", action="store_true",
                    help="mostra um aviso de 'dados fictícios' acima da agenda")
    args = ap.parse_args()

    for fluxo in (sys.stdout, sys.stderr):           # evita erro de acentos no terminal do Windows
        try:
            fluxo.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

    print(f"gerar_agenda.py · versão {VERSAO}")
    pagina = Path(args.pagina)
    if not pagina.exists():
        sys.exit(f"Não encontrei a página: {pagina}")
    saida = Path(args.saida) if args.saida else pagina

    bruto = pagina.read_bytes().decode("utf-8")
    crlf = "\r\n" in bruto
    texto = bruto.replace("\r\n", "\n")
    linhas = texto.split("\n")

    abas = [a.strip() for a in args.aba.split(",") if a.strip()] if args.aba else None
    try:
        registros, ignoradas = ler_planilha(args.planilha, abas)
        grade, meta = ler_grade(linhas)
        blocos_p, secao = ler_posters(linhas, meta)
        if not grade:
            raise ErroPlanilha("Não encontrei os horários da agenda na página (<!-- DAY n --> e slots).")
        erros, avisos, colocados, sem_horario, excluidos, colocados_p = validar(
            registros, grade, args.ordem_das_linhas, args.estrito, blocos_p, meta)
    except ErroPlanilha as e:
        print(f"\nERRO: {e}\n")
        sys.exit(1)

    if secao and not secao["antigo"]:
        tinha = {norm(t) for t in secao["temas"]}
        for b in blocos_p:
            if b["cartoes"] and norm(b["tema"]) not in tinha:
                avisos.append(f"o tema “{b['tema']}” não tinha bloco de pôsteres na página; o bloco foi criado "
                              "porque a planilha tem pôster dele (confira se o tema está certo).")

    if erros:
        print(f"\nA planilha tem {len(erros)} problema(s). Nada foi escrito:\n")
        for e in erros:
            print("  ✗", e)
        if any("a ordem" in e for e in erros):
            print("\nOrdem → horário (igual em todos os dias): " + mapa_de_ordens(grade))
        if any("pôster" in e for e in erros) and blocos_p:
            print("\nBlocos de pôsteres da página: " + mapa_de_posteres(blocos_p))
        if avisos:
            print("\nAvisos:")
            for a in avisos:
                print("  !", a)
        print()
        sys.exit(1)

    try:
        provisorio = any(r.get("provisorio") for r in list(colocados.values()) + list(colocados_p.values()))
        aviso_html = AVISO_EXEMPLO if args.exemplo else (AVISO_PROVISORIO if provisorio else None)
        novo, intro_ok, notas_p = montar_pagina(texto, linhas, grade, meta, colocados, aviso_html, colocados_p,
                                                blocos_p, secao)
    except ErroPlanilha as e:
        print(f"\nERRO: {e}\n")
        sys.exit(1)

    if saida.resolve() == pagina.resolve():
        shutil.copyfile(pagina, str(pagina) + ".bak")
    saida.write_bytes((novo.replace("\n", "\r\n") if crlf else novo).encode("utf-8"))

    # ---------------------------------------------------------------- relatório
    vagos = sorted(k for k in grade if k not in colocados)
    posters = [r for r in registros if r["tipo"] == "poster"]
    orais = [r for r in registros if r["tipo"] in ("oral", "invited")]
    print(f"\nPlanilha: {args.planilha}  ({len(registros)} apresentações lidas"
          + (f"; {ignoradas} inscrições sem apresentação foram ignoradas" if ignoradas else "") + ")")
    print(f"Agenda:   {len(colocados)} de {len(grade)} horários preenchidos")
    if vagos and not colocados:
        print("Vagos: todos os horários continuam como TBD.")
    elif vagos:
        print("Vagos (ficam como TBD): " + ", ".join(f"dia {d} {h}" for d, h in vagos))
    if posters and not blocos_p:
        print(f"Pôsteres: {len(posters)} cadastrado(s) (esta página não tem a seção de pôsteres; nada foi colocado)")
    elif posters or blocos_p:
        mostrados = [b for b in blocos_p if b["cartoes"] or b["vagas_n"]] or blocos_p
        por_tema = ", ".join(f"{b['tema']}: {len(b['cartoes'])}" for b in mostrados)
        print(f"Pôsteres: {len(colocados_p)} de {len(posters)} cadastrado(s) colocados na página ({por_tema})")
        livres = [(b["tema"], n) for b in blocos_p for n in b["vagas_n"]]
        if livres:
            print("          Slots reservados (TBD) mantidos: "
                  + ", ".join(f"{tema} P{n:02d}" for tema, n in livres))
        renum = [(n, r) for n, r in sorted(colocados_p.items())
                 if r["poster_no"] is not None and r["poster_no"] != n]
        if renum:
            print(f"          Renumerados para manter a sequência P01–P{max(colocados_p):02d} sem lacunas: "
                  + "; ".join(f"{rotulo(r)} P{r['poster_no']:02d}→P{n:02d}" for n, r in renum))
    for nota in notas_p:
        print("\n" + nota)
    if sem_horario:
        print(f"\nSem horário ainda: {len(sem_horario)} apresentação(ões) — não aparecem na página até receberem "
              "Order (posição no dia) ou Start (horário):")
        for r in sem_horario:
            print(f"  · {rotulo(r)} — {r['tipo']}, sessão {r['dia_sessao'] or '?'}: {r['titulo'][:70]}")
    if orais and (sem_horario or args.ordem_das_linhas):
        print("\nOrais por dia (pela sessão temática) x horários de oral na agenda:")
        for d in sorted(meta):
            n_o = sum(1 for r in orais if r["tipo"] == "oral" and (r["dia"] or r["dia_sessao"]) == d)
            n_s = sum(1 for (dd, _), g in grade.items() if dd == d and g["tipo"] == "oral")
            nota = "" if n_o <= n_s else f"  ← {n_o - n_s} a mais: mude o dia de alguns ou passe para pôster"
            print(f"  Dia {d} ({meta[d].get('tema', '')}): {n_o} oral(is) para {n_s} horários{nota}")
    if not intro_ok:
        avisos.append("não achei a frase de introdução da agenda; ela não foi alterada.")
    if excluidos:
        print(f"\nFICARAM DE FORA da página ({len(excluidos)}) — corrija na planilha e rode de novo:")
        for e in excluidos:
            print("  ✗", e)
        if any("a ordem" in e for e in excluidos):
            print("\n  Ordem → horário (igual em todos os dias): " + mapa_de_ordens(grade))
        if any("pôster" in e for e in excluidos) and blocos_p:
            print("\n  Blocos de pôsteres da página: " + mapa_de_posteres(blocos_p))
    if avisos:
        print(f"\nConfira na planilha ({len(avisos)}):")
        for a in avisos:
            print("  !", a)
    if provisorio:
        print("\nATENÇÃO: a ordem desta página é PROVISÓRIA (ordem das linhas). Não publique antes de definir a ordem real.")
    print(f"\nPágina escrita em: {saida}" + (f"  (cópia anterior em {pagina}.bak)" if saida.resolve() == pagina.resolve() else ""))
    print()


if __name__ == "__main__":
    main()